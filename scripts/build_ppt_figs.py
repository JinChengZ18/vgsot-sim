#!/usr/bin/env python3
"""Route clean matplotlib panels through the chapter deck so the (a)(b)(c) letters
live in the PPT, not baked into the plots (portability) — adapted from the sibling
project's eda/build_ppt_figs.py flow.

Pipeline (python-pptx + LibreOffice + PyMuPDF):
  1. panel generators (scripts/10_rtn_reservoir/plot_*_figs.py) emit letter-less
     panels to article/ppt/panels/;
  2. this script appends one slide per figure (panels in a row, (a)(b)(c) at each
     panel's top-left) to article/ppt/Chapter02_local.pptx — hand-built slides are
     never modified: our slides are tagged in the notes (AUTOFIG:) and replaced on
     re-runs; the deck is created if missing and backed up once (.bak);
  3. LibreOffice renders the deck to PDF and PyMuPDF exports each appended slide,
     CLIPPED to the panel row, to article/figs/Chapter02_local_NN.png.

Run (Windows):
  PYTHONPATH=src python scripts/10_rtn_reservoir/plot_node_figs.py
  python scripts/build_ppt_figs.py
"""
from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

import fitz
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

REPO = Path(__file__).resolve().parents[1]
PANELS = REPO / "article" / "ppt" / "panels"
PPT = REPO / "article" / "ppt"
OUT = REPO / "article" / "figs"
SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"
TAG = "AUTOFIG:"                        # marker (in slide notes) for our appended slides
MARGIN, GAP, TOP = 0.25, 0.14, 0.55     # inches, within the slide
LETTER_DX, LETTER_DY = 0.05, 0.02       # panel letter sits inside each panel's top-left

# chapter deck -> list of figures (stem in article/ppt/panels/, letters, output name)
DECKS = {
    "Chapter02_local.pptx": [
        {"stem": "ch02_19", "panels": "abc", "out": "Chapter02_local_19"},
        {"stem": "ch02_20", "panels": "abc", "out": "Chapter02_local_20"},
    ],
}


def is_auto(slide):
    return (slide.has_notes_slide
            and slide.notes_slide.notes_text_frame.text.startswith(TAG))


def remove_autoslides(prs):
    """Drop only our own previously-appended slides — remove BOTH the slide-id entry
    and its relationship so the orphaned part is not re-serialized (repeated runs
    would otherwise corrupt the package for LibreOffice)."""
    lst = prs.slides._sldIdLst
    for sid, slide in list(zip(list(lst), list(prs.slides))):
        if is_auto(slide):
            rId = sid.get(qn("r:id"))
            lst.remove(sid)
            prs.part.drop_rel(rId)


def append_fig(prs, fig):
    layout = min(prs.slide_layouts, key=lambda l: len(l.placeholders))  # blankest
    slide = prs.slides.add_slide(layout)
    for ph in list(slide.placeholders):            # make it truly blank
        ph._element.getparent().remove(ph._element)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    bg.line.fill.background(); bg.shadow.inherit = False
    letters = fig["panels"]
    paths = [PANELS / f"{fig['stem']}_{c}.png" for c in letters]
    n = len(paths)
    slide_w_in = prs.slide_width / 914400.0
    avail = slide_w_in - 2 * MARGIN - (n - 1) * GAP
    pw = avail / n
    ph_max = 0.0
    for i, path in enumerate(paths):
        im = Image.open(path)
        ph = pw * im.height / im.width
        ph_max = max(ph_max, ph)
        x = MARGIN + i * (pw + GAP)
        slide.shapes.add_picture(str(path), Inches(x), Inches(TOP), width=Inches(pw))
        if n > 1:                                  # single-panel figures get no letter
            tb = slide.shapes.add_textbox(Inches(x + LETTER_DX), Inches(TOP + LETTER_DY),
                                          Inches(0.6), Inches(0.30))
            tf = tb.text_frame; tf.word_wrap = False
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            r = tf.paragraphs[0].add_run()
            r.text = f"({letters[i]})"
            r.font.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
    slide.notes_slide.notes_text_frame.text = TAG + fig["stem"]
    clip = (MARGIN - 0.12, TOP - 0.06, slide_w_in - MARGIN + 0.12, TOP + ph_max + 0.1)
    return clip                                    # inches: (x0, y0, x1, y1)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    for deck_name, figs in DECKS.items():
        deck = PPT / deck_name
        if not deck.exists():
            prs0 = Presentation()                  # create a fresh 4:3 deck
            prs0.save(str(deck))
            print("created", deck_name)
        else:
            bak = deck.with_suffix(".pptx.bak")
            if not bak.exists():
                shutil.copy2(deck, bak)            # one-time backup of hand-built deck
        prs = Presentation(str(deck))
        remove_autoslides(prs)
        base = len(prs.slides._sldIdLst)           # hand-built slide count
        clips = [append_fig(prs, f) for f in figs]
        prs.save(str(deck))
        subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf",
                        "--outdir", str(PPT), str(deck)],
                       check=True, capture_output=True)
        pdf = fitz.open(str(PPT / deck_name.replace(".pptx", ".pdf")))
        for j, fig in enumerate(figs):
            page = pdf.load_page(base + j)
            x0, y0, x1, y1 = (v * 72 for v in clips[j])   # inches -> PDF points
            pix = page.get_pixmap(clip=fitz.Rect(x0, y0, x1, y1), dpi=230)
            pix.save(str(OUT / f"{fig['out']}.png"))
            print("wrote", (OUT / f"{fig['out']}.png").relative_to(REPO))
        pdf.close()


if __name__ == "__main__":
    main()
